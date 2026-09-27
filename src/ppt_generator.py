from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor

import textwrap
import os
from src.config import (
    MAX_RECOMMENDATION_LENGTH,
    MAX_POLICY_CLAUSE_LENGTH,
    MAX_EVIDENCE_LENGTH,
    TEXT_WRAP_WIDTH
)


# =====================================================
# Helper
# =====================================================

def wrap_text(text, width=TEXT_WRAP_WIDTH):
    if not text:
        return ""

    return "\n".join(
        textwrap.wrap(
            str(text),
            width
        )
    )


class PPTGenerator:

    def __init__(self):

        self.prs = Presentation()

        # Widescreen 16:9
        self.prs.slide_width = Inches(13.33)
        self.prs.slide_height = Inches(7.5)

    # -------------------------------------------------
    # Blue section header
    # -------------------------------------------------

    def add_section_header(
        self,
        slide,
        text,
        left,
        top,
        width
    ):

        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            left,
            top,
            width,
            Inches(0.35)
        )

        fill = shape.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(
            31,
            78,
            121
        )

        shape.line.color.rgb = RGBColor(
            31,
            78,
            121
        )

        tf = shape.text_frame
        tf.clear()

        p = tf.paragraphs[0]

        p.text = text

        p.alignment = PP_ALIGN.LEFT

        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(
            255,
            255,
            255
        )

    # -------------------------------------------------
    # Generate Presentation
    # -------------------------------------------------

    def generate(
        self,
        pitch,
        output_path
    ):
        # =====================================================
        # TITLE SLIDE
        # =====================================================

        slide = self.prs.slides.add_slide(
            self.prs.slide_layouts[6]   # Blank layout
        )

# -----------------------------
# Company Name
# -----------------------------

        title_box = slide.shapes.add_textbox(
            Inches(1),
            Inches(2.0),
            Inches(11.33),
            Inches(0.8)
        )

        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = pitch["company"]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(30)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 0, 0)

# -----------------------------
# Industry
# -----------------------------

        industry_box = slide.shapes.add_textbox(
            Inches(1),
            Inches(3.15),
            Inches(11.33),
            Inches(0.5)
        )

        tf = industry_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"Industry: {pitch['industry']}"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(90, 90, 90)

# -----------------------------
# Deck Title
# -----------------------------

        deck_box = slide.shapes.add_textbox(
           Inches(1),
           Inches(4.35),
           Inches(11.33),
           Inches(1.2)
        )

        tf = deck_box.text_frame

        p = tf.paragraphs[0]
        p.text = "Insurance Recommendation Deck"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(24)
        p.font.color.rgb = RGBColor(80, 80, 80)

        p = tf.add_paragraph()
        p.text = "Generated using RAG + LLM"
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(120, 120, 120)
        # =====================================================
        # EXECUTIVE SUMMARY
        # =====================================================

        slide = self.prs.slides.add_slide(
            self.prs.slide_layouts[5]
        )

        slide.shapes.title.text = "Executive Summary"

        title = slide.shapes.title.text_frame.paragraphs[0]
        title.alignment = PP_ALIGN.CENTER
        title.font.size = Pt(28)
        title.font.bold = True
        title.font.color.rgb = RGBColor(0, 0, 0)

        rows = len(pitch["slides"]) + 1
        cols = 2

        table = slide.shapes.add_table(
            rows,
            cols,
            Inches(0.7),
            Inches(1.45),   # moved down
            Inches(11.5),
            Inches(3.2)
        ).table

        table.columns[0].width = Inches(5.7)
        table.columns[1].width = Inches(5.8)

        table.cell(0, 0).text = "Business Risk"
        table.cell(0, 1).text = "Recommended Policy"

        for c in range(cols):

            cell = table.cell(0, c)

            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(
                31,
                78,
                121
            )

            cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

            cell.text_frame.paragraphs[0].font.bold = True
            cell.text_frame.paragraphs[0].font.size = Pt(16)
            cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(
                255,
                255,
                255
            )

        for i, s in enumerate(
            pitch["slides"],
            start=1
        ):

            table.cell(i, 0).text = s["title"]
            table.cell(i, 1).text = s["recommended_policy"]

            for j in range(cols):

                para = table.cell(i, j).text_frame.paragraphs[0]

                para.font.size = Pt(14)
                para.font.color.rgb = RGBColor(0, 0, 0)

        # =====================================================
        # RISK SLIDES
        # =====================================================

        for slide_data in pitch["slides"]:

            slide = self.prs.slides.add_slide(
                self.prs.slide_layouts[5]
            )

            title = slide.shapes.title
            title.left = Inches(0.5)
            title.top = Inches(0.1)
            title.width = Inches(12.3)
            title.height = Inches(0.55)

            title.text = slide_data["title"]

            p = title.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(30)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0, 0, 0)

            current_top = 1.1
                        # =====================================================
            # Recommended Policy
            # =====================================================

            self.add_section_header(
                slide,
                "Recommended Policy",
                Inches(0.5),
                Inches(current_top),
                Inches(5.8)
            )

            current_top += 0.45

            box = slide.shapes.add_textbox(
                Inches(0.6),
                Inches(current_top),
                Inches(5.8),
                Inches(0.55)
            )

            tf = box.text_frame
            tf.word_wrap = True

            p = tf.paragraphs[0]
            p.text = slide_data["recommended_policy"]
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0, 0, 0)

            current_top += 0.75

            # =====================================================
            # Recommendation
            # =====================================================

            self.add_section_header(
                slide,
                "Recommendation",
                Inches(0.5),
                Inches(current_top),
                Inches(5.8)
            )

            current_top += 0.45

            rec_box = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                Inches(0.55),
                Inches(current_top),
                Inches(5.8),
                Inches(1.25)
            )

            rec_box.fill.solid()
            rec_box.fill.fore_color.rgb = RGBColor(242, 242, 242)
            rec_box.line.color.rgb = RGBColor(200, 200, 200)

            tf = rec_box.text_frame
            tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

            p = tf.paragraphs[0]
            recommendation = slide_data["recommendation"]

            if len(recommendation) > MAX_RECOMMENDATION_LENGTH:
                recommendation = (
                    recommendation[:MAX_RECOMMENDATION_LENGTH] + "..."
                )

            p.text = wrap_text(
                recommendation,
                TEXT_WRAP_WIDTH
            )
            p.font.size = Pt(14)
            p.font.color.rgb = RGBColor(0, 0, 0)

            current_top += 1.50

            # =====================================================
            # Key Benefits
            # =====================================================

            self.add_section_header(
                slide,
                "Key Benefits",
                Inches(0.5),
                Inches(current_top),
                Inches(5.8)
            )

            current_top += 0.45

            benefit_box = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                Inches(0.55),
                Inches(current_top),
                Inches(5.8),
                Inches(2.0)
            )

            benefit_box.fill.solid()
            benefit_box.fill.fore_color.rgb = RGBColor(250, 250, 250)
            benefit_box.line.color.rgb = RGBColor(200, 200, 200)

            tf = benefit_box.text_frame
            tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

            benefit = slide_data["key_benefit"]

            if len(benefit) > MAX_POLICY_CLAUSE_LENGTH:
                benefit = (
                    benefit[:MAX_POLICY_CLAUSE_LENGTH] + "..."
                )
            
            benefit_lines = [
                x.strip()
                for x in benefit.split(".")
                if x.strip()
            ]

            if not benefit_lines:
                benefit_lines = [benefit]
            first = True

            for line in benefit_lines[:6]:

                if first:
                    para = tf.paragraphs[0]
                    first = False
                else:
                    para = tf.add_paragraph()

                para.text = "• " + line
                para.font.size = Pt(14)
                para.font.color.rgb = RGBColor(0, 0, 0)
                            # =====================================================
            # Supporting Evidence
            # =====================================================

            self.add_section_header(
                slide,
                "Supporting Evidence",
                Inches(7.0),
                Inches(0.8),
                Inches(5.6)
            )

            evidence_top = 1.25

            for ev in slide_data["evidence"][:2]:

                card = slide.shapes.add_shape(
                    MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                    Inches(7.0),
                    Inches(evidence_top),
                    Inches(5.8),
                    Inches(2.0)
                )

                card.fill.solid()
                card.fill.fore_color.rgb = RGBColor(250, 250, 250)
                card.line.color.rgb = RGBColor(180, 180, 180)

                tf = card.text_frame
                tf.word_wrap = True
                tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

                # Heading
                p = tf.paragraphs[0]
                p.text = (
                    f"{ev['policy_name']}  |  "
                    f"Page {ev['page_number']}"
                )
                p.font.bold = True
                p.font.size = Pt(13)
                p.font.color.rgb = RGBColor(0, 0, 0)

                # Evidence snippet
                snippet = ev["text"].replace("\n", " ")

                if len(snippet) > MAX_EVIDENCE_LENGTH:
                    snippet = (
                        snippet[:MAX_EVIDENCE_LENGTH] + "..."
                    )

                p = tf.add_paragraph()
                p.text = wrap_text(
                            snippet,
                            TEXT_WRAP_WIDTH
                        )
                p.font.size = Pt(11)
                p.font.color.rgb = RGBColor(70, 70, 70)

                evidence_top += 2.15

            # =====================================================
            # Footer
            # =====================================================

            footer = slide.shapes.add_textbox(
                Inches(0.4),
                Inches(7.0),
                Inches(12.5),
                Inches(0.25)
            )

            tf = footer.text_frame

            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.text = (
                "Source: Retrieved Insurance Policy Documents | "
                "Generated using Retrieval-Augmented Generation (RAG)"
            )
            p.font.size = Pt(9)
            p.font.color.rgb = RGBColor(120, 120, 120)

        # =====================================================
        # SAVE PRESENTATION
        # =====================================================

        os.makedirs(
            os.path.dirname(output_path),
            exist_ok=True
        )

        self.prs.save(output_path)

        print("\nPresentation saved to:")
        print(output_path)

        return output_path


# =====================================================
# Public Function
# =====================================================

def create_pitch_deck(
    pitch,
    output_path
):

    generator = PPTGenerator()

    return generator.generate(
        pitch,
        output_path
    )